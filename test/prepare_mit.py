"""Convert the original HFC or Rental Excel files into TurnPoint's standard CSV format.

Run from the repository root:
    python test/prepare_mit.py "data/MIT AI negotiation competition" data/formatted
Requires openpyxl. Matching wide outcome workbooks are read: these
contain both dialogue and outcomes. Other exercises and duplicate long-form exports are ignored.
Use --exercise rental for Rental; --conditions optionally selects conditions.
"""

import argparse
import os
import sys
import ast
from collections import Counter
from contextlib import ExitStack
import csv
from pathlib import Path

# EDIT THESE SETTINGS when opening this script and clicking Run.
# Relative paths start at the repository root. Generated files are replaced.
SOURCE = "data/MIT AI negotiation competition"
DESTINATION = "data/formatted"
EXERCISE = "hfc"  # "hfc" or "rental"
CONDITIONS = None  # None = all conditions; or e.g. ["round2"]

import re

from openpyxl import load_workbook

DIMENSIONS = ("instrumental", "self", "process", "relationship")
ROLE = re.compile(r"^\s*(consultant|coo|landlord|tenant)\s*:\s*", re.IGNORECASE)
PRIVATE = re.compile(r"<negotiation_(preparation|strategy)>.*?</negotiation_\1>", re.DOTALL | re.IGNORECASE)


def prepare_mit(source, destination, conditions=None, exercise="hfc"):
    """Read selected MIT source workbooks and write turns and linked outcome tables.

    source is the supplied MIT folder containing outcomes/. Reads all supplied
    exercise conditions unless condition names are given. Participant IDs are the
    source roles (consultant/coo or landlord/tenant); agent IDs remain metadata. Conversation IDs
    combine condition, exercise, and original Excel row, avoiding collisions.

    Removes appended SVI questionnaire responses and tagged preparation/strategy
    blocks from dialogue. Retains original message positions as turn numbers.
    Missing outcome cells stay missing.
    Unparseable/truncated dialogue, unidentified speakers, and conversations
    without both roles are excluded and logged in excluded_conversations.csv.
    Source workbooks are streamed and never modified. Output files are overwritten.
    """
    source, destination = Path(source), Path(destination)
    files = sorted((source / "outcomes").glob(f"*_{exercise}_wide_outcomes_deid.xlsx"))
    available = {p.name.removesuffix(f"_{exercise}_wide_outcomes_deid.xlsx") for p in files}
    if conditions:
        if set(conditions) - available:
            raise ValueError(f"Conditions not found: {sorted(set(conditions) - available)}")
        files = [p for p in files if p.name.removesuffix(f"_{exercise}_wide_outcomes_deid.xlsx") in conditions]
    if not files:
        raise FileNotFoundError(f"No matching wide outcome workbooks in {source / 'outcomes'}")
    destination.mkdir(parents=True, exist_ok=True)
    fields = {
        "turns": ["conversation_id", "turn_number", "speaker_id", "text", "agent_id", "condition", "source_file", "source_row"],
        "conversation_outcomes": ["conversation_id", "value_created"],
        "individual_outcomes": ["conversation_id", "speaker_id", *DIMENSIONS],
        "excluded_conversations": ["source_file", "source_row", "reason"],
    }
    required = {"conversation", "role_1", "role_2", "value_created"} | {
        f"svi_{dimension}_{side}" for dimension in DIMENSIONS for side in (1, 2)}
    report = []
    with ExitStack() as stack:
        writers = {}
        for name, columns in fields.items():
            handle = stack.enter_context((destination / f"{name}.csv").open("w", newline="", encoding="utf-8"))
            writers[name] = csv.DictWriter(handle, fieldnames=columns)
            writers[name].writeheader()
        for path in files:
            counts = Counter()
            condition = path.name.removesuffix(f"_{exercise}_wide_outcomes_deid.xlsx")
            workbook = load_workbook(path, read_only=True, data_only=True)
            try:
                rows = workbook.active.iter_rows(values_only=True)
                headers = next(rows)
                if required - set(headers):
                    raise ValueError(f"{path.name}: missing columns {sorted(required - set(headers))}")
                for source_row, values in enumerate(rows, 2):
                    row = dict(zip(headers, values))
                    roles = [str(row[f"role_{side}"]).strip().lower() for side in (1, 2)]
                    if set(roles) != ({"consultant", "coo"} if exercise == "hfc" else {"landlord", "tenant"}):
                        raise ValueError(f"{path.name}, row {source_row}: unexpected participant roles")
                    reason, cleaned = None, []
                    removed_surveys = removed_private = 0
                    try:
                        messages = ast.literal_eval(row["conversation"])
                        if not isinstance(messages, list) or not messages or not all(isinstance(t, str) for t in messages):
                            raise ValueError("Expected a nonempty list of messages")
                    except (SyntaxError, ValueError, TypeError):
                        messages, reason = [], "unparseable_or_truncated_dialogue"
                    survey_started = False
                    for number, message in enumerate(messages, 1):
                        message, removed = PRIVATE.subn("", message)
                        removed_private += removed
                        message = message.replace("\\n", "\n").strip()
                        if removed and not message:
                            continue
                        if re.search(r"</?negotiation_(preparation|strategy)", message, re.IGNORECASE):
                            reason = "incomplete_private_block"
                            break
                        match = ROLE.match(message)
                        if not match:
                            reason = "unidentified_speaker"
                            break
                        text = message[match.end():].strip()
                        # SVI responses are numbered 1..16 and contain agreement ratings.
                        survey = (re.search(r"(?m)^\s*1[.)]\s", text)
                                  and re.search(r"(?m)^\s*16[.)]\s", text)
                                  and re.search(r"\b(?:agree|disagree)\b", text, re.IGNORECASE))
                        if survey:
                            survey_started = True
                            removed_surveys += 1
                            continue
                        if survey_started or not text:
                            reason = "unexpected_text_after_survey" if survey_started else "empty_turn"
                            break
                        cleaned.append((number, match.group(1).lower(), text))
                    if reason is None and {speaker for _, speaker, _ in cleaned} != set(roles):
                        reason = "missing_participant_dialogue"
                    if reason:
                        writers["excluded_conversations"].writerow({"source_file": path.name, "source_row": source_row, "reason": reason})
                        counts[reason] += 1
                        continue
                    conversation_id = f"{condition}:{exercise}:{source_row}"
                    for number, speaker, text in cleaned:
                        side = roles.index(speaker) + 1
                        agent = row.get(f"name_deid_{side}", row.get(f"name_{side}"))
                        writers["turns"].writerow({"conversation_id": conversation_id, "turn_number": number,
                            "speaker_id": speaker, "text": text, "agent_id": agent,
                            "condition": condition, "source_file": path.name, "source_row": source_row})
                    writers["conversation_outcomes"].writerow({"conversation_id": conversation_id, "value_created": row["value_created"]})
                    for side, speaker in enumerate(roles, 1):
                        writers["individual_outcomes"].writerow({"conversation_id": conversation_id, "speaker_id": speaker,
                            **{dimension: row[f"svi_{dimension}_{side}"] for dimension in DIMENSIONS}})
                    counts["conversations"] += 1
                    counts["turns"] += len(cleaned)
                    counts["survey_messages_removed"] += removed_surveys
                    counts["private_blocks_removed"] += removed_private
            finally:
                workbook.close()
            report.append({"source_file": path.name, **dict(counts)})
            print(f"{path.name}: {dict(counts)}", flush=True)
    import json
    (destination / "preparation_summary.json").write_text(json.dumps({"exercise": exercise, "files": report}, indent=2))


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("source", type=Path)
    parser.add_argument("destination", type=Path)
    parser.add_argument("--exercise", choices=["hfc", "rental"], default="hfc", help="MIT exercise (default: hfc)")
    parser.add_argument("--conditions", nargs="+", help="Condition names; default: all supplied conditions")
    arguments = sys.argv[1:]
    if not arguments:
        os.chdir(Path(__file__).resolve().parents[1])
        arguments = [str(SOURCE), str(DESTINATION), "--exercise", EXERCISE]
        if CONDITIONS:
            arguments += ["--conditions", *CONDITIONS]
    args = parser.parse_args(arguments)
    prepare_mit(args.source, args.destination, args.conditions, args.exercise)
